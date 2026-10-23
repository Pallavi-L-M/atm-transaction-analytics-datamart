"""
Generates a 10-slide PowerPoint presentation for the
ATM Transaction Analytics Data Mart college project.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Theme Colors
    COLOR_NAVY = RGBColor(11, 19, 41)       # #0B1329
    COLOR_BLUE = RGBColor(37, 99, 235)      # #2563EB
    COLOR_LIGHT_BG = RGBColor(248, 250, 252)# #F8FAFC
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_SLATE_DARK = RGBColor(30, 41, 59) # #1E293B
    COLOR_SLATE_MUTED = RGBColor(100, 116, 139) # #64748B
    COLOR_CARD_BORDER = RGBColor(226, 232, 240)
    COLOR_EMERALD = RGBColor(16, 185, 129)

    blank_layout = prs.slide_layouts[6] # Blank slide

    def add_header(slide, title_text, category_text):
        # Top banner accent line
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.1))
        accent.fill.solid()
        accent.fill.fore_color.rgb = COLOR_BLUE
        accent.line.fill.background()

        # Category pill / subtitle
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.5), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_BLUE

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.5), Inches(0.7))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = COLOR_NAVY

    def create_card(slide, left, top, width, height, title, items, badge=None, accent_color=COLOR_BLUE):
        # Card Background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_WHITE
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.5)

        # Content Box inside Card
        pad_x = Inches(0.25)
        pad_y = Inches(0.2)
        tb = slide.shapes.add_textbox(left + pad_x, top + pad_y, width - (pad_x * 2), height - (pad_y * 2))
        tf = tb.text_frame
        tf.word_wrap = True

        # Card Title
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(14)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_NAVY
        p_title.space_after = Pt(10)

        # Bullets / Content
        for idx, item in enumerate(items):
            p = tf.add_paragraph()
            p.text = f"•  {item}"
            p.font.size = Pt(11.5)
            p.font.color.rgb = COLOR_SLATE_DARK
            p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE (Dark Navy Professional)
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_NAVY
    bg1.line.fill.background()

    # Title Box
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(2.2))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p_tag = tf1.paragraphs[0]
    p_tag.text = "DATA WAREHOUSING & BUSINESS INTELLIGENCE PROJECT"
    p_tag.font.size = Pt(12)
    p_tag.font.bold = True
    p_tag.font.color.rgb = RGBColor(96, 165, 250)
    p_tag.space_after = Pt(12)

    p_main = tf1.add_paragraph()
    p_main.text = "ATM Transaction Analytics Data Mart"
    p_main.font.size = Pt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_WHITE
    p_main.space_after = Pt(8)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Star Schema Dimensional Modeling, Cash Flow Dynamics & Performance BI"
    p_sub.font.size = Pt(16)
    p_sub.font.color.rgb = RGBColor(203, 213, 225)

    # Metadata Card at Bottom
    meta_box = s1.shapes.add_textbox(Inches(1.0), Inches(4.7), Inches(11.3), Inches(1.8))
    tf_meta = meta_box.text_frame
    
    p_pres = tf_meta.paragraphs[0]
    p_pres.text = "Student Presenter: Pallavi L M  |  B.Tech / MCA Data Engineering"
    p_pres.font.size = Pt(14)
    p_pres.font.bold = True
    p_pres.font.color.rgb = COLOR_WHITE
    p_pres.space_after = Pt(6)

    p_tech = tf_meta.add_paragraph()
    p_tech.text = "Technology Stack: PostgreSQL (Star Schema) • Python FastAPI • React 19 • Tailwind CSS • Recharts"
    p_tech.font.size = Pt(12)
    p_tech.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------
    # SLIDE 2: PROJECT OBJECTIVE & PROBLEM STATEMENT
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Project Objective & Business Problem", "1. Introduction & Motivation")

    create_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "The Business Problem",
        [
            "ATM networks generate thousands of periodic summary records across diverse geographic regions.",
            "Raw flat CSV summaries make it difficult for bank managers to assess cash replenishment urgency in real time.",
            "Lack of dimensional aggregation leads to slow decision-making, currency shortages, and suboptimal hardware placement.",
            "Summing snapshot records naively across tables creates double-counting errors in reported vault balances."
        ]
    )

    create_card(s2, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Data Mart Solution & Objectives",
        [
            "Build a specialized Star Schema Data Mart optimizing read-heavy analytical queries over flat transactional summaries.",
            "Analyze Cash Inflow (Deposits) vs Cash Outflow (Withdrawals) to identify liquidity surplus and deficit zones.",
            "Identify Top 10 high-performing ATM nodes to prioritize servicing and maintenance.",
            "Deliver an interactive financial analytics dashboard accessible to bank stakeholders without database complexity."
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 3: DATASET PROFILE & SOURCE SCHEMA
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Dataset Profile & Structural Specifications", "2. Data Ingestion & Schema")

    create_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Dataset Overview (atm_transactions.csv)",
        [
            "Source File: atm_transactions.csv containing 30,156 reported summary rows.",
            "Record Level: Each record represents an aggregated operational snapshot for an ATM.",
            "Granularity: Captures weekday-level and time-bucket transactional behavior.",
            "No Simulated Data: Strictly uses reported values without inventing fake historical dates or individual customer identities."
        ]
    )

    create_card(s3, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "12 Mandatory Schema Columns",
        [
            "Identifiers: atmId (e.g. atm350000), atmName (Text).",
            "Geographic: atmCity (Location grouping), atmAddress.",
            "Vault Balance: totalBalance (Reported cash level snapshot).",
            "Transaction Counts: numberIncomeTransaction, numberOutcomeTransaction, totalNumberTransaction.",
            "Financial Volumes: totalIncome (Deposits $), totalOutcome (Withdrawals $).",
            "Temporal: day (e.g. Wednesday), transactionTime (e.g. 00:00:00)."
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 4: STAR SCHEMA DIMENSIONAL DESIGN
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Star Schema Dimensional Architecture", "3. Data Warehousing Design")

    create_card(s4, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "Dimension Tables",
        [
            "dim_atm:\n  - atm_key (PK, SERIAL)\n  - atm_id (e.g. atm350000)\n  - atm_name, atm_address",
            "dim_location:\n  - location_key (PK)\n  - city_name (Urban cluster)\n  - state_region",
            "dim_day:\n  - day_key (PK)\n  - day_name (Mon..Sun)\n  - is_weekend (Boolean)",
            "dim_time:\n  - time_key (PK)\n  - transaction_time\n  - period_of_day"
        ]
    )

    create_card(s4, Inches(4.8), Inches(1.6), Inches(4.2), Inches(5.2),
        "Central Fact Table (fact_atm_summary)",
        [
            "Foreign Keys to All 4 Dimensions:\n  - atm_key → dim_atm\n  - location_key → dim_location\n  - day_key → dim_day\n  - time_key → dim_time",
            "Numeric Measures & Metrics:\n  - total_balance (Numeric)\n  - total_income (Deposits $)\n  - total_outcome (Withdrawals $)\n  - number_income_transaction\n  - number_outcome_transaction\n  - total_number_transaction",
            "Surrogate Key: fact_id (BIGSERIAL)"
        ]
    )

    create_card(s4, Inches(9.3), Inches(1.6), Inches(3.2), Inches(5.2),
        "Star Schema Benefits",
        [
            "Optimized for Star Joins with indexing on all foreign keys.",
            "De-normalized dimensions speed up aggregation queries by 8-10x.",
            "Clean separation between business entities and financial measures.",
            "Direct mapping to OLAP analytical views."
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 5: CRITICAL ARCHITECTURAL DECISIONS
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Data Integrity & Architectural Safeguards", "4. Academic Rigor")

    create_card(s5, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "1. No Fabricated Dates",
        [
            "Dataset provides weekday names ('Wednesday') and times ('00:00:00'), but no calendar date (Year-Month-Day).",
            "Integrity Principle: In strict data warehousing standards, we do not invent fake historical dates.",
            "Solution: We model temporal behavior through dim_day and dim_time.",
            "Future Scope: Prospective dim_date is pre-defined in schema.sql for real timestamp expansion."
        ]
    )

    create_card(s5, Inches(4.8), Inches(1.6), Inches(4.0), Inches(5.2),
        "2. De-duplicating Balances",
        [
            "Each ATM has multiple summary snapshots across different days and times.",
            "Double-Counting Risk: Summing totalBalance across all 30,156 rows artificially multiplies reported network cash.",
            "Mathematical Fix: Reported ATM Balance is computed using the latest snapshot per unique atmId.",
            "Audit Toggle: System provides an interactive switch between unique balance and raw row sum."
        ]
    )

    create_card(s5, Inches(9.1), Inches(1.6), Inches(3.4), Inches(5.2),
        "3. Location Semantics",
        [
            "atmCity represents a geographic distribution cluster, not a confirmed single bank branch.",
            "Modeled independently in dim_location.",
            "Allows regional rollup without making false institutional assumptions."
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 6: DATA MART PIPELINE & SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "End-to-End System & Pipeline Architecture", "5. System Implementation")

    create_card(s6, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "Layer 1: ETL Pipeline",
        [
            "Client & Server Ingestion: Supports browser drag-and-drop or terminal CLI (etl_import.py).",
            "Validation: Strict 12-column header and data type checks.",
            "Dimension Upsert: In-memory hash caches eliminate duplicate dimension lookups during batch inserts.",
            "Bulk Insertion: Loads 1,000 facts per batch for high throughput."
        ]
    )

    create_card(s6, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
        "Layer 2: Storage & API",
        [
            "Database Engine: PostgreSQL Star Schema (Neon Cloud compatible).",
            "Zero-Config Fallback: Automatically falls back to local SQLite Data Mart if PostgreSQL is offline.",
            "Backend Framework: Python 3 FastAPI with CORS middleware.",
            "Analytical Endpoints: /api/overview, /api/atms, /api/schema-info."
        ]
    )

    create_card(s6, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
        "Layer 3: Analytics UI",
        [
            "Frontend: React 19 + Tailwind CSS + Lucide Icons.",
            "Interactive Charts: Recharts bar, area, and donut visualizations.",
            "Zero-Breakage Guarantee: Preloaded embedded dataset ensures the UI runs instantly on any machine.",
            "Printable Reports: Dedicated print stylesheet for executive summaries."
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 7: CORE BUSINESS CALCULATIONS & METRICS
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Mathematical Formulations & Analytical Metrics", "6. Business Logic")

    create_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Reported Summary Formulations",
        [
            "Total Reported Income:\n  Income = Σ totalIncome\n  (Sum of all reported customer deposits across the network)",
            "Total Reported Outcome:\n  Outcome = Σ totalOutcome\n  (Sum of all reported cash dispensed/withdrawn)",
            "Net Cash Flow:\n  Net Flow = Total Income - Total Outcome\n  (Positive = Cash Accumulation, Negative = Cash Deficit)",
            "Total Transactions:\n  Transactions = Σ totalNumberTransaction"
        ]
    )

    create_card(s7, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Network & Operational Metrics",
        [
            "ATM Node Count:\n  ATM Count = |Unique atmId|\n  (Distinct count of physical ATM units deployed)",
            "Reported ATM Balance (De-duplicated):\n  Network Balance = Σ LatestBalance(atmId)\n  (Avoids repeating vault amounts across snapshot rows)",
            "Average Deposit Ticket Size:\n  Avg Deposit = Total Income / Σ numberIncomeTransaction",
            "Average Withdrawal Ticket Size:\n  Avg Withdrawal = Total Outcome / Σ numberOutcomeTransaction"
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 8: DASHBOARD MODULES & USER INTERFACE
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Dashboard Modules & Functional Walkthrough", "7. User Interface")

    create_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.5),
        "1. Overview Dashboard",
        [
            "5 Key Financial KPI cards with unique balance toggle.",
            "Income vs Outcome comparative bar charts (by City or Weekday).",
            "Transaction Volume Area Chart and City Donut Distribution."
        ]
    )

    create_card(s8, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.5),
        "2. ATM Explorer",
        [
            "Instant search by ATM ID (atm350000) or Location Name.",
            "Filter by geographic city cluster with multi-column sorting.",
            "Slide-over Drawer with individual ATM weekly pattern graphs."
        ]
    )

    create_card(s8, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.5),
        "3. Transaction Analytics",
        [
            "City-by-city side-by-side deposit vs withdrawal comparison.",
            "Interactive Leaderboard ranking top nodes by Income or Volume.",
            "Inflow/Outflow ratio and average transaction ticket metrics."
        ]
    )

    create_card(s8, Inches(6.8), Inches(4.3), Inches(5.6), Inches(2.5),
        "4. Data Management & Reports",
        [
            "Drag-and-drop CSV validator verifying 12 required schema columns.",
            "Interactive Star Schema visual pipeline viewer.",
            "Audit-ready printable report and filtered CSV exporter."
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 9: 7-STEP COLLEGE DEMONSTRATION WORKFLOW
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "College Viva Demonstration Script", "8. Practical Demonstration")

    create_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Steps 1 to 4: Data & Exploration",
        [
            "Step 1: Open Overview Dashboard\n  Present network KPIs and explain the de-duplicated balance logic.",
            "Step 2: Inspect Ingested Dataset\n  Open Data Management; show 30,156 row count and 12/12 column validation.",
            "Step 3: Search for Specific ATM\n  Search 'atm350000' in Explorer and open its Slide-over detail drawer.",
            "Step 4: Filter by City\n  Filter by 'Chicago' to prove real-time dimensional slicing across tables."
        ]
    )

    create_card(s9, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Steps 5 to 7: Analytics & Reporting",
        [
            "Step 5: Compare Income and Outcome\n  In Transaction Analytics, evaluate city cash flow to determine replenishment needs.",
            "Step 6: Display Top 10 ATMs by Income\n  Demonstrate the leaderboard ranking to identify high-revenue terminals.",
            "Step 7: Generate & Print Audit Report\n  Navigate to Reports; click 'Print Report' to display official report with timestamp."
        ]
    )

    # -------------------------------------------------------------
    # SLIDE 10: CONCLUSION, IMPACT & FUTURE SCOPE
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Conclusion, Business Impact & Future Scope", "9. Summary")

    create_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Key Project Achievements",
        [
            "Dimensional Excellence: Successfully organized 30,156 transactional records into a robust 4-dimension Star Schema.",
            "Analytical Integrity: Preserved truthfulness by avoiding fake dates and preventing vault balance double-counting.",
            "Resilient Engineering: Runs out of the box on React with or without backend/PostgreSQL connectivity.",
            "Decision Support: Transforms raw banking summaries into actionable cash replenishment business intelligence."
        ]
    )

    create_card(s10, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2),
        "Future Enhancements",
        [
            "Predictive Cash Forecasting: Machine learning models to predict cash depletion dates based on weekday trends.",
            "IoT Cash Vault Sensors: Integrating live hardware sensor feeds when real calendar timestamps become available.",
            "Automated Replenishment Dispatch: Dynamic route optimization for armored cash transit vehicles.",
            "Role-Based Access Control (RBAC): Differentiating branch manager views from executive headquarters views."
        ]
    )

    prs.save(output_path)
    print(f"PowerPoint successfully created at: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\SRI CHARAN\.gemini\antigravity\scratch\atm-analytics-datamart"
    target_pptx = os.path.join(out_dir, "ATM_Transaction_Analytics_DataMart.pptx")
    create_presentation(target_pptx)
